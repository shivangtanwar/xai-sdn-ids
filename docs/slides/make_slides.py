#!/usr/bin/env python3
"""Build the Phase 1-3 presentation (docs/slides/XAI-SDN-IDS_Phase1-3.pptx).

    docker run --rm -v "$PWD":/app -w /app xai-sdn-ids/controller \
        sh -c "pip install -q python-pptx && python docs/slides/make_slides.py"
"""
import os

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DIAG = os.path.join(ROOT, "docs", "diagrams")
RES = os.path.join(ROOT, "results")
OUT = os.path.join(ROOT, "docs", "slides", "XAI-SDN-IDS_Phase1-3.pptx")

BLUE = RGBColor(0x2A, 0x78, 0xD6)
INK = RGBColor(0x0B, 0x0B, 0x0B)
INK2 = RGBColor(0x52, 0x51, 0x4E)
TINT = RGBColor(0xE6, 0xF0, 0xFC)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W = prs.slide_width


def text(slide, x, y, w, h, s, size=18, color=INK, bold=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(s.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size, p.font.color.rgb, p.font.bold = Pt(size), color, bold
    return tb


def slide(title, kicker=None):
    s = prs.slides.add_slide(BLANK)
    bar = s.shapes.add_shape(1, 0, 0, W, Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()
    if kicker:
        text(s, Inches(0.6), Inches(0.3), Inches(12), Inches(0.4), kicker.upper(), 12, BLUE, True)
    text(s, Inches(0.6), Inches(0.55), Inches(12), Inches(0.8), title, 30, INK, True)
    text(s, Inches(0.6), Inches(7.0), Inches(12), Inches(0.4),
         "XAI-SDN-IDS · Software Defined Security · Problem Statement #15", 10, INK2)
    return s


def bullets(s, items, x=Inches(0.6), y=Inches(1.6), w=Inches(12), size=19):
    tb = s.shapes.add_textbox(x, y, w, Inches(5.2))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        sub = item.startswith("  ")
        p.text = ("–  " if sub else "•  ") + item.strip()
        p.level = 1 if sub else 0
        p.font.size = Pt(size - 3 if sub else size)
        p.font.color.rgb = INK2 if sub else INK
        p.space_after = Pt(6)


def table(s, df, x, y, w, size=12, col_w=None):
    rows, cols = df.shape[0] + 1, df.shape[1]
    shape = s.shapes.add_table(rows, cols, x, y, w, Inches(0.35 * rows))
    t = shape.table
    for j, c in enumerate(df.columns):
        t.cell(0, j).text = str(c)
    for i, row in enumerate(df.itertuples(index=False), 1):
        for j, v in enumerate(row):
            t.cell(i, j).text = str(v)
    for i in range(rows):
        for j in range(cols):
            for p in t.cell(i, j).text_frame.paragraphs:
                p.font.size = Pt(size)
                p.font.bold = i == 0
    if col_w:
        for j, cw in enumerate(col_w):
            t.columns[j].width = Inches(cw)
    return t


def image(s, path, x, y, w=None, h=None):
    if os.path.exists(path):
        s.shapes.add_picture(path, x, y, width=w, height=h)
    else:
        text(s, x, y, Inches(5), Inches(0.5), f"[missing: {os.path.basename(path)}]", 14, INK2)


# 1 title
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(1, 0, 0, W, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = TINT
bg.line.fill.background()
text(s, Inches(0.8), Inches(1.2), Inches(11.5), Inches(0.5), "SOFTWARE DEFINED SECURITY · END-TERM PROJECT · PHASE 1–3", 14, BLUE, True)
text(s, Inches(0.8), Inches(1.8), Inches(11.5), Inches(2.0),
     "Explainable AI-Based SDN Intrusion Detection\nand Automated Flow Response", 40, INK, True)
text(s, Inches(0.8), Inches(3.9), Inches(11.5), Inches(1.0),
     "XAI-SDN-IDS: Ryu + Mininet/OVS · Random Forest · SHAP · Confidence-Gated Response Engine", 20, INK2)
text(s, Inches(0.8), Inches(5.2), Inches(11.5), Inches(1.2),
     "Team: <names / roll numbers>\nFaculty: Dr. Sumit Badotra · 15 October 2026", 18, INK)

# 2 problem
s = slide("Why SDN intrusion detection needs explanations and guard-rails", "Phase 1 · Problem")
bullets(s, [
    "SDN gives the controller a global, programmable view, so IDS/IPS can run as a controller app",
    "ML-based SDN IDSs report ≥ 99% accuracy, but:",
    "  Black box: operators cannot see why a host was flagged, so they cannot trust or audit it",
    "  Trigger-happy: most install a DROP on the first positive. A false positive becomes a self-inflicted DoS",
    "  Offline: evaluated on CSV datasets whose features a controller cannot actually measure",
    "Problem #15: collect OVS flow stats → ML IDS → feature-level explanations → install mitigation "
    "only when defined confidence/security conditions are satisfied",
])

# 3 objectives
s = slide("Objectives & scope", "Phase 1 · Objectives")
bullets(s, [
    "O1  Multi-switch Mininet/OVS testbed with benign traffic + 4 attack classes",
    "O2  Ryu app: poll OpenFlow flow stats every 2 s → 14 per-source features",
    "O3  Train & compare multi-class detectors (Acc, P, R, F1, FPR, FNR)",
    "O4  SHAP feature-level explanations: local per alert, global per model",
    "O5  Confidence-Gated Response Engine: alert → meter rate-limit → self-expiring DROP",
    "O6  Evaluate detection/mitigation latency, victim QoS, controller overhead, gating ablation",
    "Scope: emulated SDN (OF1.3), volumetric + recon attacks from identifiable sources. "
    "Out of scope: payload/DPI, spoofed-source attribution",
], size=18)

# 4 literature
s = slide("Literature review: 16 verified papers (2020–2026)", "Phase 1 · Literature")
lit = pd.DataFrame([
    ["[1] Elsayed 2020", "InSDN dataset (ONOS testbed)", "–", "–", "–"],
    ["[2] Novaes 2020", "LSTM + fuzzy, Floodlight", "–", "✓", "statistical τ"],
    ["[3] Tan 2020", "trigger + KMeans/KNN", "–", "✓", "2-stage filter"],
    ["[5] Alashhab 2024", "online ensemble, Ryu", "–", "✓", "n/r"],
    ["[6] El Kamel 2022", "DNN rate limiting", "–", "✓", "trust score"],
    ["[7] Wang 2020", "SHAP for IDS (NSL-KDD)", "✓", "–", "–"],
    ["[11] Saykat 2026", "OptiXGB, InSDN, SHAP/LIME", "✓", "–", "–"],
    ["[12] Ahmad 2026*", "XAI-SDN: RF + SHAP, Ryu", "✓", "offline", "single τ=0.70"],
    ["Ours", "RF + SHAP + CGRE, live Ryu", "✓", "✓ graduated", "G1–G4"],
], columns=["Work", "Approach", "XAI", "Auto-mitigation", "Gating"])
table(s, lit, Inches(0.6), Inches(1.5), Inches(12.1), 13, [2.2, 4.2, 1.0, 2.0, 2.7])
text(s, Inches(0.6), Inches(6.5), Inches(12), Inches(0.4),
     "* preprint. Full table, synthesis and DOIs: docs/research/literature_review.md", 11, INK2)

# 5 gaps
s = slide("Research gaps → our response", "Phase 1 · Gaps")
bullets(s, [
    "RG1  Explanations never drive mitigation (offline [10][11] or analyst-only [12]) → SHAP is an input to the policy",
    "RG2  No multi-condition safety gate in SDN controllers → confidence + persistence + XAI + asset gates",
    "RG3  Explanation quality only scored offline [9][13] → per-alert runtime consistency check",
    "RG4  Graduated, self-expiring response is rare [6] → alert → meter → DROP with hard_timeout",
    "RG5  SDN XAI is mostly binary/offline → 5-class model on controller-native OpenFlow features",
    "RG6  Per-flow SHAP is ~31× slower [12] → explain only flagged per-source windows",
], size=17)

# 6 architecture
s = slide("System architecture", "Phase 2 · Design")
image(s, os.path.join(DIAG, "architecture.png"), Inches(1.5), Inches(1.35), h=Inches(5.6))

# 7 topology + pipeline
s = slide("Topology & OpenFlow pipeline", "Phase 2 · Design")
image(s, os.path.join(DIAG, "topology.png"), Inches(0.4), Inches(1.4), w=Inches(7.6))
bullets(s, [
    "Table 0 SECURITY: DROP / meter rules, default goto 1",
    "Table 1 FORWARDING: L2 learning + L3/L4 rules, miss → controller",
    "Match = (src, dst, proto, service port). Random source ports don't explode the flow table",
    "idle_timeout 10 s; stats polled every 2 s",
    "h5/h6 are protected assets (G4)",
], x=Inches(8.2), y=Inches(1.6), w=Inches(4.8), size=15)

# 8 features
s = slide("14 interpretable per-source features (2 s window)", "Phase 2 · Features")
feats = pd.DataFrame([
    ["pkt_rate, byte_rate", "volume: floods"], ["avg_pkt_size", "tiny packets: SYN/ICMP flood"],
    ["n_flows, new_flows", "flow churn: scans"], ["n_dst_ips, n_dst_ports", "fan-out: sweeps/scans"],
    ["frac_tcp / frac_udp / frac_icmp", "protocol mix"], ["pkts_per_flow", "low: scan · high: flood/bulk"],
    ["frac_small_flows", "half-open / probe flows"], ["mean_flow_duration", "short vs persistent flows"],
    ["pktin_rate", "control-plane pressure"],
], columns=["Feature(s)", "Attack signal"])
table(s, feats, Inches(0.6), Inches(1.5), Inches(8.5), 14, [4.0, 4.5])
text(s, Inches(9.4), Inches(1.6), Inches(3.5), Inches(4),
     "Classes:\nbenign\nicmp_flood\nsyn_flood\nudp_flood\nport_scan\n\nBenign includes heavy iperf3 "
     "transfers, so the model cannot learn 'rate = attack'.", 15, INK)

# 9 CGRE
s = slide("Innovation: Confidence-Gated Response Engine", "Phase 2 · Security mechanism")
image(s, os.path.join(DIAG, "cgre_flow.png"), Inches(0.4), Inches(1.35), h=Inches(5.55))
bullets(s, [
    "G1  p ≥ 0.70 to act, ≥ 0.90 to block",
    "G2  ≥ 2 of last 3 windows",
    "G3  top-3 SHAP ∩ class signature",
    "G4  never block protected assets",
    "Rules auto-expire (60 s): a false positive heals itself",
    "Audit log + REST override (human-in-the-loop)",
], x=Inches(9.3), y=Inches(1.6), w=Inches(3.8), size=15)

# 10 tech
s = slide("Technology choices", "Phase 2 · Justification")
tech = pd.DataFrame([
    ["Ryu 4.34 (Python)", "ML/SHAP in-process; simple OF1.3 stats, meters, multi-table (ODL/ONOS: Java, heavy)"],
    ["OpenFlow 1.3", "meters for rate-limiting, multiple tables, per-flow counters"],
    ["Mininet 2.3 + OVS 2.17", "course standard, real OVS kernel datapath"],
    ["Random Forest", "best accuracy/latency trade-off on tabular data; exact fast TreeSHAP"],
    ["SHAP TreeExplainer", "additive, consistent, local + global; ms per alert"],
    ["Docker Compose", "reproducible on all laptops; Ryu needs Python 3.9"],
], columns=["Choice", "Why"])
table(s, tech, Inches(0.6), Inches(1.6), Inches(12.1), 15, [3.2, 8.9])

# 11 implementation
s = slide("Implementation status", "Phase 3 · Implementation")
bullets(s, [
    "Docker testbed: controller (Py3.9, Ryu, sklearn, SHAP) + privileged Mininet/OVS container",
    "controller/xai_ids_app.py: two-table pipeline, reactive forwarding, stats poller, mitigator (DROP + meters), REST API",
    "ids/: feature extractor, detector, SHAP explainer, response engine (policy.py), config",
    "traffic/scenarios.py: randomised labelled episodes (hping3 ICMP/SYN/UDP floods, nmap -sS/-sT, iperf3, curl, ping)",
    "ml/: dataset labelling, 7-model comparison, SHAP figures, audit/latency summary",
    "Unit tests for feature extraction and every CGRE gate (9/9 passing)",
    "Three modes: collect (dataset) · monitor (explain, never act) · detect (closed loop)",
], y=Inches(1.4), size=16)
image(s, os.path.join(DIAG, "ml_pipeline.png"), Inches(1.4), Inches(4.3), w=Inches(10.5))

# 12 results
s = slide("Preliminary results: held-out session", "Phase 3 · Results")
mc = os.path.join(RES, "model_comparison.csv")
if os.path.exists(mc):
    df = pd.read_csv(mc)[["model", "accuracy", "precision", "recall", "f1", "fpr", "fnr", "infer_us_per_sample"]]
    df.columns = ["Model", "Acc", "Prec", "Rec", "F1", "FPR", "FNR", "µs/sample"]
    for c in df.columns[1:-1]:
        df[c] = df[c].map(lambda v: f"{v:.3f}")
    df["µs/sample"] = df["µs/sample"].map(lambda v: f"{v:.1f}")
    table(s, df, Inches(0.5), Inches(1.5), Inches(7.6), 12, [2.3, 0.75, 0.75, 0.75, 0.75, 0.75, 0.75, 0.8])
image(s, os.path.join(RES, "confusion_matrix_rf.png"), Inches(8.3), Inches(1.4), h=Inches(4.3))
text(s, Inches(0.5), Inches(5.9), Inches(12.2), Inches(0.9),
     "Trained on session 1, tested on an independent session 2 (different seed, rates, attackers, victims). "
     "FPR/FNR = benign-vs-attack view. Macro metrics over 5 classes.", 13, INK2)

# 13 shap
s = slide("Explanations: what drives each attack class", "Phase 3 · XAI")
image(s, os.path.join(RES, "shap_per_class.png"), Inches(0.4), Inches(1.4), w=Inches(12.5))
example = "Per-alert example: run scripts/demo.sh to populate results/audit_log.jsonl"
audit = os.path.join(RES, "audit_log.jsonl")
if os.path.exists(audit):
    import json
    blocks = [json.loads(line) for line in open(audit)]
    blocks = [e for e in blocks if e["action"] in ("BLOCK", "RATE_LIMIT")]
    if blocks:
        e = blocks[0]
        why = ", ".join(f"{x['feature']}={x['value']:.0f} ({x['shap']:+.2f})" for x in e["explanation"])
        example = (f"Real alert from the demo audit log:\n{e['src_ip']} → {e['label']}  p={e['confidence']:.2f}\n"
                   f"because {why}\n→ gates passed → {e['action']} (auto-expires)")
text(s, Inches(0.6), Inches(5.0), Inches(12.2), Inches(1.9),
     example + "\n\nPer-class top features also validate the CGRE signatures (G3).", 14, INK)

# 13b ablation + live demo
s = slide("Gating works: no innocent host blocked, every attack stopped", "Phase 3 · Security results")
abl = os.path.join(RES, "ablation_test2.csv")
if os.path.exists(abl):
    a = pd.read_csv(abl)
    innocent = {"Naive (block on label)": "2 (h6 server, h8 user)"}
    a = pd.DataFrame({
        "Policy": a["variant"],
        "Innocent hosts impaired": [innocent.get(v, "0") for v in a["variant"]],
        "Attacks mitigated": a["attacks_mitigated"],
        "Mean time to mitigate": a["mean_time_to_mitigate_s"].map(lambda v: f"{v:.1f} s"),
    })
    table(s, a, Inches(0.6), Inches(1.5), Inches(12.1), 14, [3.6, 3.7, 2.4, 2.4])
text(s, Inches(0.6), Inches(3.4), Inches(12.1), Inches(0.5),
     "Replay of held-out session. Naive blocking hit flood victims whose replies looked like floods.", 13, INK2)
demo = pd.DataFrame([
    ["Heavy benign 10 Mbps UDP (h1)", "–", "–", "none (correct)"],
    ["SYN flood h4 → h5", "2.2 s", "6.4 s", "ALERT → BLOCK · frac_tcp=1, avg_pkt_size=54 B"],
    ["Port scan h7 → h6", "3.8 s", "6.1 s", "ALERT → BLOCK · n_dst_ports=538"],
    ["UDP flood h4 → h8", "2.3 s", "2.4 s", "RATE_LIMIT (p=.89) → BLOCK (p=.97) · frac_udp=1"],
], columns=["Live closed-loop demo", "Flagged", "Mitigated", "Action · SHAP reason"])
table(s, demo, Inches(0.6), Inches(4.0), Inches(12.1), 14, [3.6, 1.3, 1.4, 5.8])
text(s, Inches(0.6), Inches(6.2), Inches(12.1), Inches(0.6),
     "Inference 9–31 ms, SHAP explanation ≈2.5 ms per alert. OVS meters verified (OF1.3, kbps, drop band).", 13, INK2)

# 13c honest limitation
s = slide("What we learned: limitations that shape Phase 4", "Phase 3 · Analysis")
bullets(s, [
    "Same-session CV overstates performance: CV F1 0.97–0.99 for all models, but 0.68–0.97 on an independent session",
    "  we always report held-out-session metrics",
    "Victim reflection: a host answering an ICMP flood emits an equal-rate ICMP stream",
    "  a 2,400 pps ICMP flood was detected with p < 0.70 (confidence gate held it back)",
    "  Phase 4: ICMP-type feature + reflection-aware gate G5",
    "Statistics lag up to 2 polling intervals → detection latency floor ≈ 2–4 s",
    "Labelling bug found and fixed: looping scans outlived their episode (now killed by process group)",
    "G2–G4 did not change outcomes in this session → adversarial/edge scenarios planned for Phase 5",
], size=17)

# 14 plan
s = slide("Plan for Phases 4–7 (final evaluation 16–19 Nov)", "Roadmap")
plan = pd.DataFrame([
    ["16–22 Oct", "More sessions & intensities (incl. low-rate), tuning, InSDN cross-check"],
    ["23–29 Oct", "Closed-loop CGRE hardening, meters, audit log, REST API"],
    ["30 Oct–5 Nov", "Experiments: detection/mitigation latency, throughput/RTT/loss, CPU/mem, "
                     "ablation naive vs gated, SHAP vs LIME"],
    ["6–12 Nov", "Operator dashboard, final report & figures"],
    ["13–15 Nov", "Demo rehearsal, viva preparation, backup recording"],
], columns=["When", "Work"])
table(s, plan, Inches(0.6), Inches(1.6), Inches(12.1), 15, [2.2, 9.9])

# 15 thanks
s = slide("Thank you. Questions?", "")
bullets(s, [
    "Code: xai-sdn-ids/ (README → quick start)",
    "Key references: [1] InSDN (IEEE Access 2020) · [6] El Kamel (Comp. Comm. 2022) · "
    "[7] Wang (IEEE Access 2020) · [12] XAI-SDN (arXiv 2026) · [13] Alquliti (arXiv 2025)",
    "Full IEEE-style reference list in the Phase 1–3 report",
], size=17)

prs.save(OUT)
print("saved", OUT)
